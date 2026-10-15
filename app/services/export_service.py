import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.models import Transaction


HEADERS = ["id", "tipo", "categoria", "descricao", "contraparte", "valor", "emissao", "vencimento", "pagamento", "status", "forma_pagamento"]


def _rows(transactions: list[Transaction]):
    for t in transactions:
        yield [t.id, t.transaction_type, t.category, t.description, t.counterparty, float(t.amount),
               t.issue_date.isoformat(), t.due_date.isoformat(), t.payment_date.isoformat() if t.payment_date else "",
               t.status, t.payment_method or ""]


def export_csv(scenario_id: int, transactions: list[Transaction]) -> Path:
    settings.exports_dir.mkdir(parents=True, exist_ok=True)
    path = settings.exports_dir / f"cenario-{scenario_id}-movimentacoes.csv"
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file, delimiter=";")
        writer.writerow(["SIMULAÇÃO EDUCACIONAL — USO EXCLUSIVO PARA FINS PEDAGÓGICOS"])
        writer.writerow(HEADERS)
        writer.writerows(_rows(transactions))
    return path


def export_xlsx(scenario_id: int, transactions: list[Transaction]) -> Path:
    settings.exports_dir.mkdir(parents=True, exist_ok=True)
    path = settings.exports_dir / f"cenario-{scenario_id}-movimentacoes.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Movimentações"
    sheet.append(["SIMULAÇÃO EDUCACIONAL — USO EXCLUSIVO PARA FINS PEDAGÓGICOS"])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    sheet["A1"].font = Font(bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="B03A2E")
    sheet.append(HEADERS)
    for cell in sheet[2]:
        cell.font = Font(bold=True)
    for row in _rows(transactions):
        sheet.append(row)
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:K{sheet.max_row}"
    for index, column in enumerate(sheet.columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = min(max(len(str(c.value or "")) for c in column) + 2, 45)
    workbook.save(path)
    return path
